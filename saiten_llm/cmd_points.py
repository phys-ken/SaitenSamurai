"""cmd_points.py — 配点の検討と、確定した配点への換算（S5）。

  stats       全員の正誤から統計と配点案を出す（先生が選ぶ・直す材料）
  set-points  先生が承認した配点を書き込み、仮の2点満点を本当の点数に換算する
              （正→満点 / 誤→0 / 部分→先生が決めた部分点。決めていなければ
               未採点に戻し、先生が個別に点を入れる）
"""
import json
import math
import shutil
from datetime import datetime
from pathlib import Path

from saiten_llm import allocate
from saiten_llm import judgments as jmod
from saiten_llm.engine import call, open_bridge
from saiten_llm.workspace import now_iso

ALLOCATION_FILE = "配点の検討.json"


def _say(*lines):
    for line in lines:
        print(line, flush=True)


def _desc_questions(bridge):
    return (bridge.state.get("descriptive") or {}).get("questions") or []


def build_matrix(bridge):
    """生徒×項目の表（0〜1）。Returns: (項目の一覧, 生徒の一覧, 表)"""
    from scoring_engine import load_mark2_results, load_template, score_answers
    mode = bridge.state["app_mode"]
    fmt = bridge.state["mark_format"]
    items, rows, students = [], [], []
    mark_results = {}
    if mode != "descriptive_only":
        template = load_template(bridge.state["answer_key"], mark_format=fmt)
        for s in load_mark2_results(bridge.state["omr_result"],
                                    skip_questions=bridge.state["skip_questions"]):
            res = score_answers(s["answers"], template, mark_format=fmt)
            mark_results[s["image"]] = {q: 1.0 if r["correct"] else 0.0
                                        for q, r in res["results"].items()}
        for q in sorted(template):
            items.append({"kind": "mark", "id": str(q),
                          "label": f"問{template[q].get('group_label', q)}",
                          "points": template[q]["配点"]})
    desc = _desc_questions(bridge) if mode != "mark_only" else []
    for q in desc:
        items.append({"kind": "desc", "id": q["id"], "label": f"{q['id']}（{q['name']}）",
                      "points": q["max_score"]})
    names = sorted(mark_results) if mark_results else sorted(bridge.list_sheet_files()["files"])
    scores = bridge._desc_scores["scores"]
    for name in names:
        row = []
        for item in items:
            if item["kind"] == "mark":
                row.append(mark_results.get(name, {}).get(int(item["id"]), 0.0))
            else:
                q = next(d for d in desc if d["id"] == item["id"])
                v = (scores.get(name) or {}).get(q["id"])
                row.append(0.0 if v is None else v / q["max_score"])
        students.append(name)
        rows.append(row)
    return items, students, rows


def _unconfirmed(project, bridge):
    """配点を決める前に揃っていないもの"""
    missing = []
    data = jmod.load(project)
    for q in _desc_questions(bridge) if bridge.state["app_mode"] != "mark_only" else []:
        if not (data["questions"].get(q["id"]) or {}).get("confirmed"):
            missing.append(f"記述 {q['id']}（{q['name']}）の先生の確定")
    if bridge.state["app_mode"] != "descriptive_only":
        if not (project.data.get("answer_key") or {}).get("teacher_confirmed_at"):
            missing.append("正答の先生の確定（confirm answer-key）")
        if not (project.data.get("marks_check") or {}).get("teacher_confirmed_at"):
            missing.append("マーク確認の先生の確定（confirm marks）")
    return missing


def cmd_stats(project, args):
    bridge, _ = open_bridge(project, write=False)
    missing = _unconfirmed(project, bridge)
    if missing:
        _say("まだ確定していないものがあります。統計は参考として出しますが、配点は決められません:")
        for m in missing:
            _say(f"  ・{m}")
    items, students, matrix = build_matrix(bridge)
    if len(students) < 3:
        _say("生徒が少なすぎて統計を出せません。")
        return 1
    stats = allocate.item_stats(matrix)
    plans = allocate.plans(matrix, stats, args.total)
    _say(f"生徒 {len(students)} 人・{len(items)} 項目"
         "（記述は 得点÷満点 で計算。配点を決める前の部分は 1/2 として計算）",
         "", "項目        正答率  識別力  上下差")
    for item, s in zip(items, stats):
        r = "  —  " if s["r"] is None else f"{s['r']:+.2f}"
        _say(f"  {item['label']:<10} {s['p']:5.2f}  {r}  {s['ul']:+.2f}")
    negative = [it["label"] for it, s in zip(items, stats) if s["r"] is not None and s["r"] < 0]
    if negative:
        _say("", "⚠ 識別力がマイナスの項目（成績上位の生徒ほど間違えている）: " + "、".join(negative),
             "  出題や正答の誤りの疑いがあります。配点を考える前に先生に伝えてください。")
    if len(students) < 30:
        _say(f"  （生徒が {len(students)} 人と少ないので、識別力は偶然で大きくぶれます）")
    _say("", f"配点案（満点 {args.total} 点・各項目1点以上の整数。どれを選ぶか・どう直すかは先生が決めます）",
         "  満点が項目数で割り切れないときは、端数の1点を重みの端数が大きい項目から順に足しています。")
    for p in plans:
        alpha = "—" if p["alpha"] is None else f"{p['alpha']:.2f}"
        _say(f"  {p['label']}: 平均 {p['mean']:.1f} / 標準偏差 {p['sd']:.1f} / α {alpha} / "
             f"最低 {p['min']:.0f}・最高 {p['max']:.0f}",
             "    配点: " + ", ".join(f"{it['label']}={pt}" for it, pt in zip(items, p["points"])),
             "    分布（満点の10%刻み）: " + " ".join(map(str, p["histogram_10"])))
    out = project.root / ALLOCATION_FILE
    out.write_text(json.dumps({
        "made_at": now_iso(), "total": args.total, "students": len(students),
        "items": [{**it, **s} for it, s in zip(items, stats)], "plans": plans,
        "negative_discrimination": negative, "unconfirmed": missing},
        ensure_ascii=False, indent=1), encoding="utf-8")
    _say("", f"詳しい数値: {out}")
    return 0


def _write_mark_points(key_path, points):
    """answer_key.xlsx の配点列だけを書き換える（書式はそのまま。前の版は控える）"""
    import openpyxl
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    shutil.copy2(key_path, key_path.with_name(f"{key_path.stem}.backup-{stamp}{key_path.suffix}"))
    wb = openpyxl.load_workbook(key_path)
    ws = wb.active
    header = [c.value for c in ws[1]]
    q_col, p_col = header.index("問題番号") + 1, header.index("配点") + 1
    written = set()
    for row in range(2, ws.max_row + 1):
        q = ws.cell(row, q_col).value
        key = str(q).split("-")[0].strip() if q is not None else None
        if key in points:
            ws.cell(row, p_col).value = int(points[key])
            written.add(key)
    wb.save(key_path)
    missing = set(points) - written
    if missing:
        raise ValueError(f"answer_key.xlsx に無い問題番号があります: {', '.join(sorted(missing))}")


def partial_points(desc, partial, ratio):
    """部分の最終点を問題ごとに決める。

    Args:
        desc: {問題ID: 満点}
        partial: {問題ID: 部分の点}（問題ごとの指定。ratio より優先）
        ratio: 全問共通の割合（満点 × 割合を四捨五入）。None なら使わない

    Returns:
        {問題ID: 部分の点}（指定の無い問題は含めない＝部分は未採点に戻す）

    Raises:
        ValueError: 部分の点が 1 以上・満点未満の整数にならないとき
    """
    unknown = set(partial) - set(desc)
    if unknown:
        raise ValueError("partial に書いた問題が descriptive にありません: "
                         + ", ".join(sorted(unknown)))
    out = {}
    for qid, new_max in desc.items():
        if qid in partial:
            value = partial[qid]
            if isinstance(value, bool) or not isinstance(value, int):
                raise ValueError(f"{qid} の部分の点は整数にしてください: {value!r}")
        elif ratio is not None:
            value = math.floor(new_max * ratio + 0.5)
        else:
            continue
        if not 1 <= value < new_max:
            raise ValueError(f"{qid} の部分の点が {value} 点になります。"
                             f"1 点以上・満点（{new_max} 点）未満にしてください。")
        out[qid] = value
    return out


def cmd_set_points(project, args):
    if not args.teacher_approved:
        _say("配点は先生が決めます。先生の承認を得てから --teacher-approved を付けて実行してください。")
        return 2
    bridge, adapter = open_bridge(project)
    missing = _unconfirmed(project, bridge)
    if missing:
        _say("まだ確定していないものがあるので、配点を書き込めません:")
        for m in missing:
            _say(f"  ・{m}")
        return 2
    # Windows の PowerShell は BOM 付きで書くことがあるので utf-8-sig で読む
    spec = json.loads(Path(args.input).read_text(encoding="utf-8-sig"))
    # "marks": "keep" は、先生が answer_key.xlsx に入れた配点をそのまま使う
    keep_marks = spec.get("marks") == "keep"
    marks = {} if keep_marks else {str(k): int(v) for k, v in (spec.get("marks") or {}).items()}
    desc = {str(k): int(v) for k, v in (spec.get("descriptive") or {}).items()}
    if any(v <= 0 for v in list(marks.values()) + list(desc.values())):
        _say("配点は1以上の整数にしてください。")
        return 1
    questions = {q["id"]: q for q in _desc_questions(bridge)}
    unknown = set(desc) - set(questions)
    if unknown:
        _say(f"記述問題が見つかりません: {', '.join(sorted(unknown))}")
        return 1
    ratio = getattr(args, "partial_ratio", None)
    if ratio is not None and not 0 < ratio < 1:
        _say("--partial-ratio は 0 より大きく 1 より小さい値にしてください（例: 0.5）。")
        return 1
    try:
        partial = partial_points(desc, {str(k): v for k, v in (spec.get("partial") or {}).items()},
                                 ratio)
    except ValueError as e:
        _say(str(e))
        return 1

    if marks:
        key_path = Path(bridge.state["answer_key"])
        _write_mark_points(key_path, marks)
        adapter.file_returns = [str(key_path)]
        call(bridge, "select_answer_key")
        ks = bridge.state["key_summary"]
        if not ks["ok"]:
            _say("配点を書き込んだ answer_key.xlsx にエラーがあります:", *ks["errors"])
            return 1
        project.data["answer_key"]["hash"] = jmod.file_hash(key_path)
        project.data["answer_key"]["points_set_at"] = now_iso()
        _say(f"マークの配点を書き込みました: {ks['stats_line']}")
    elif keep_marks and bridge.state["app_mode"] != "descriptive_only":
        project.data["answer_key"]["points_set_at"] = now_iso()
        _say(f"マークの配点は answer_key.xlsx のまま使います: "
             f"{(bridge.state.get('key_summary') or {}).get('stats_line')}")

    data = jmod.load(project)
    files = bridge.list_sheet_files()["files"]
    for qid, new_max in desc.items():
        q = questions[qid]
        if q["max_score"] != jmod.PROVISIONAL_MAX:
            _say(f"{qid} はすでに {q['max_score']} 点満点に換算済みです（やり直さない）。")
            continue
        before = {f: (s or {}).get(qid) for f, s in bridge._desc_scores["scores"].items()}
        record = data["questions"][qid]
        # 確定を引き継げるのは、確定の後で得点が変わっていないときだけ
        confirmed = record.get("confirmed") or {}
        still_confirmed = (confirmed.get("max_score") == jmod.PROVISIONAL_MAX and confirmed.get("hash")
                           == jmod.scores_hash({f: before.get(f) for f in files}))
        part = partial.get(qid)
        call(bridge, "update_descriptive_question", qid, q["name"], new_max, q.get("aspect", 1))
        converted = {f: (new_max if v == 2 else 0 if v == 0 else part if v == 1 else None)
                     for f, v in before.items() if v is not None}
        partials = sorted(f for f, v in before.items() if v == 1)
        unscored = sorted(f for f, v in converted.items() if v is None)
        call(bridge, "set_descriptive_scores_bulk", qid, converted)
        # 換算をやり直しても（2点満点のまま換算したときなど）、最初の確定は残す
        record.setdefault("provisional_confirmed", record.get("confirmed"))
        record["converted"] = {"at": now_iso(), "max_score": new_max, "partials": partials,
                               "partial_score": part}
        if unscored or not still_confirmed:
            record["confirmed"] = None
        else:
            scores = {f: (bridge._desc_scores["scores"].get(f) or {}).get(qid) for f in files}
            record["confirmed"] = {"at": now_iso(), "hash": jmod.scores_hash(scores),
                                   "max_score": new_max,
                                   "by": "確定済みの正誤・部分を換算" if partials else "確定済みの正誤を換算"}
        rule = f"正→{new_max} / " + (f"部分→{part} / " if part is not None else "") + "誤→0"
        _say(f"{qid}（{q['name']}）を {new_max} 点満点に換算しました（{rule}）。")
        if partials and part is None:
            _say(f"  部分 {len(partials)} 人は未採点に戻しました。先生が採点侍で「未採点」に"
                 "絞り込んで点を入れ、その後 confirm question で確定してください。")
        elif unscored:
            _say(f"  2点満点の 0/1/2 以外の得点だった {len(unscored)} 人は未採点に戻しました。"
                 "先生に点を入れてもらい、confirm question で確定してください。")
        elif not still_confirmed:
            _say("  ⚠ 先生の確定の後で得点が変わっていたので、確定は引き継いでいません。"
                 "先生に確かめてもらい、confirm question で確定し直してください。")
        else:
            _say("  先生の確定は換算した点数のまま引き継ぎました。")
    jmod.save(project, data)
    project.data["points"] = {"set_at": now_iso(), "marks": "keep" if keep_marks else marks,
                              "descriptive": desc}
    if partial:
        project.data["points"]["partial"] = partial
    project.save()
    call(bridge, "save_session")
    return 0
